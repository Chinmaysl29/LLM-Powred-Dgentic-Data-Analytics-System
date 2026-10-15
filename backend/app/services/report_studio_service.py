"""Source-backed report generation and atomic scoped artifact manifests."""
from pathlib import Path
from datetime import datetime, timezone
from uuid import uuid4
import json
import os
from backend.reports.report_generator import ReportGenerator

REPORT_STUDIO_STORAGE_ROOT = Path("storage/reports/studio")
class ReportType:
    EXECUTIVE = "Executive Report"
    BOARD = "Board Report"
    INVESTOR = "Investor Report"
    WEEKLY_OPS = "Weekly Operations Report"
    MONTHLY_OPS = "Monthly Operations Report"

class ReportStudioService:
    def __init__(self, storage_dir=None):
        self.storage_dir = Path(storage_dir or REPORT_STUDIO_STORAGE_ROOT)
        self.storage_dir.mkdir(parents=True, exist_ok=True)

    def generate_report(self, title, report_type=ReportType.EXECUTIVE, df=None, workspace_id=None,
                        executive_summary=None, key_metrics=None, recommendations=None, **kwargs):
        provenance = kwargs.get("provenance")
        if not provenance or not provenance.get("analysis_run_id"):
            raise ValueError("A completed source-backed analysis run is required")
        data = kwargs.get("report_data")
        if not isinstance(data, dict): raise ValueError("Computed report content is required")
        format_type = kwargs.get("format_type", "pdf")
        artifact = ReportGenerator(self.storage_dir).generate(report_type, format_type, title, data)
        source = Path(artifact["file_path"])
        self.validate_artifact(source, artifact["format"])
        result = {"id": artifact["report_id"], "workspace_id": str(workspace_id), "title": title,
                  "report_type": report_type, "created_at": datetime.now(timezone.utc).isoformat(),
                  "format": artifact["format"], "status": "GENERATED", "provenance": provenance,
                  "summary": data.get("summary", ""), "metrics": data.get("kpis", {}),
                  "filename": source.name, "size_bytes": source.stat().st_size}
        target = self.storage_dir / f"{result['id']}.json"
        temporary = target.with_suffix(f".{uuid4().hex}.tmp")
        try:
            with temporary.open("w", encoding="utf-8") as handle:
                json.dump(result, handle, ensure_ascii=False, allow_nan=False)
                handle.flush(); os.fsync(handle.fileno())
            os.replace(temporary, target)
        finally: temporary.unlink(missing_ok=True)
        return result

    @staticmethod
    def validate_artifact(path, format_type):
        if format_type == "pdf":
            from pypdf import PdfReader
            if not PdfReader(path).pages: raise ValueError("Empty PDF artifact")
        elif format_type == "docx":
            from docx import Document
            if not Document(path).paragraphs: raise ValueError("Empty document artifact")
        elif format_type == "xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True); wb.close()
        elif format_type == "pptx":
            from pptx import Presentation
            if not Presentation(path).slides: raise ValueError("Empty presentation artifact")
        else: raise ValueError("Unsupported artifact format")

    def list_reports(self, workspace_id=None):
        reports = []
        for path in self.storage_dir.glob("rep_*.json"):
            record = json.loads(path.read_text(encoding="utf-8"))
            if not workspace_id or record.get("workspace_id") == str(workspace_id): reports.append(record)
        return reports

    def get_report(self, report_id):
        return next((r for r in self.list_reports() if r["id"] == report_id), None)

_report_studio_service = ReportStudioService()
def get_report_studio_service(): return _report_studio_service
